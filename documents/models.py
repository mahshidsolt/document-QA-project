from django.db import models
import docx
from .utils import chunk_text, store_chunks, delete_document_chunks

class Document(models.Model):
    title = models.CharField(max_length=255)
    file = models.FileField(upload_to='documents/')
    extracted_text = models.TextField(blank=True)
    uploaded_at = models.DateTimeField(auto_now_add=True)

    def __str__(self):
        return self.title

    def save(self, *args, **kwargs):
        is_new = self.pk is None
        file_changed = is_new

        if self.pk:
            old_document = Document.objects.filter(pk=self.pk).first()

            if old_document and old_document.file.name != self.file.name:
                file_changed = True

        super().save(*args, **kwargs)

        if self.file and self.file.name.lower().endswith('.docx') and file_changed:

            # اگر سند ویرایش شده، chunkهای نسخه قبلی پاک شوند
            if not is_new:
                delete_document_chunks(self.id)

            doc = docx.Document(self.file.path)

            full_text = []
            for paragraph in doc.paragraphs:
                full_text.append(paragraph.text)

            self.extracted_text = '\n'.join(full_text)

            super().save(update_fields=['extracted_text'])

            chunks = chunk_text(self.extracted_text)
            store_chunks(self.id, chunks)


class QAHistory(models.Model):
    question = models.TextField()
    answer = models.TextField()
    created_at = models.DateTimeField(auto_now_add=True)
    documents = models.ManyToManyField(
    Document,
    blank=True,
    related_name='qa_history'
    )

    def __str__(self):
        return self.question[:50]